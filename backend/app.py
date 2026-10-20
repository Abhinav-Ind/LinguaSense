import os
from datetime import datetime, timedelta
from functools import wraps

from dotenv import load_dotenv
from flask import Flask, jsonify, request
from io import BytesIO
from pypdf import PdfReader
from docx import Document
from flask_cors import CORS
from flask_jwt_extended import (
    JWTManager, create_access_token, get_jwt_identity,
    jwt_required
)
from flask_sqlalchemy import SQLAlchemy
from groq import Groq
from werkzeug.security import check_password_hash, generate_password_hash

load_dotenv()
app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("FLASK_SECRET_KEY", "change-this-secret")
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv(
    "DATABASE_URL", "sqlite:///linguasense.db"
)
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 6 * 1024 * 1024  # 6 MB total request cap
app.config["JWT_SECRET_KEY"] = os.getenv("JWT_SECRET_KEY", "change-this-jwt-secret")
app.config["JWT_ACCESS_TOKEN_EXPIRES"] = timedelta(hours=12)

CORS(app, resources={r"/api/*": {"origins": os.getenv("FRONTEND_ORIGIN", "*")}})
db = SQLAlchemy(app)
jwt = JWTManager(app)

MODEL_NAME = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
groq_client = Groq(api_key=os.getenv("GROQ_API_KEY")) if os.getenv("GROQ_API_KEY") else None

PERSONAS = {
    "Professional": "Use a polished, concise, professional tone. Preserve intent and important details.",
    "Student": "Explain clearly and naturally for a student. Prefer accessible wording without being childish.",
    "Creative": "Use expressive, engaging language while preserving the author's meaning.",
    "Friendly": "Use a warm, approachable, conversational tone.",
    "Academic": "Use precise, formal, structured language suitable for academic work."
}

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(255), unique=True, nullable=False, index=True)
    password_hash = db.Column(db.String(255), nullable=False)
    persona = db.Column(db.String(40), nullable=False, default="Professional")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class History(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    kind = db.Column(db.String(30), nullable=False)
    title = db.Column(db.String(180), nullable=False)
    input_text = db.Column(db.Text, nullable=False)
    output_text = db.Column(db.Text, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

def fail(message, status=400):
    return jsonify({"error": message}), status

def get_ai():
    if not groq_client:
        raise RuntimeError("GROQ_API_KEY is not configured. Add it to backend/.env.")
    return groq_client

def ai_text(system_prompt, user_prompt, temperature=0.35):
    client = get_ai()
    result = client.chat.completions.create(
        model=MODEL_NAME,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=temperature,
    )
    return (result.choices[0].message.content or "").strip()

def save_history(user_id, kind, title, input_text, output_text):
    row = History(
        user_id=user_id, kind=kind, title=title[:180],
        input_text=input_text, output_text=output_text
    )
    db.session.add(row)
    db.session.commit()
    return row

@app.get("/api/health")
def health():
    return jsonify({"status": "ok", "app": "LinguaSense", "model": MODEL_NAME,
                    "groq_configured": bool(os.getenv("GROQ_API_KEY"))})

@app.post("/api/auth/register")
def register():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip()
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))
    if not name or not email or not password:
        return fail("Name, email and password are required.")
    if len(password) < 8:
        return fail("Password must contain at least 8 characters.")
    if "@" not in email or "." not in email.rsplit("@", 1)[-1]:
        return fail("Enter a valid email address.")
    if User.query.filter_by(email=email).first():
        return fail("An account with this email already exists.", 409)
    user = User(name=name, email=email, password_hash=generate_password_hash(password))
    db.session.add(user)
    db.session.commit()
    token = create_access_token(identity=str(user.id))
    return jsonify({"token": token, "user": user_payload(user)}), 201

@app.post("/api/auth/login")
def login():
    data = request.get_json(silent=True) or {}
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))
    user = User.query.filter_by(email=email).first()
    if not user or not check_password_hash(user.password_hash, password):
        return fail("Incorrect email or password.", 401)
    token = create_access_token(identity=str(user.id))
    return jsonify({"token": token, "user": user_payload(user)})

def user_payload(user):
    return {"id": user.id, "name": user.name, "email": user.email,
            "persona": user.persona}

def current_user():
    return db.session.get(User, int(get_jwt_identity()))

@app.get("/api/auth/me")
@jwt_required()
def me():
    user = current_user()
    if not user:
        return fail("User not found.", 404)
    return jsonify({"user": user_payload(user)})

@app.get("/api/personas")
@jwt_required()
def personas():
    user = current_user()
    return jsonify({"personas": list(PERSONAS.keys()), "selected": user.persona})

@app.put("/api/persona")
@jwt_required()
def set_persona():
    data = request.get_json(silent=True) or {}
    persona = str(data.get("persona", "")).strip()
    if persona not in PERSONAS:
        return fail("Choose a supported persona.")
    user = current_user()
    user.persona = persona
    db.session.commit()
    return jsonify({"message": "Persona updated.", "persona": persona})

@app.post("/api/refine")
@jwt_required()
def refine():
    data = request.get_json(silent=True) or {}
    source = str(data.get("text", "")).strip()
    instruction = str(data.get("instruction", "Correct grammar, improve clarity, and preserve meaning.")).strip()
    if not source:
        return fail("Enter text to refine.")
    if len(source) > 30000:
        return fail("Text is too long. Keep it under 30,000 characters.")
    user = current_user()
    system = (
        "You are LinguaSense, an expert writing assistant. "
        f"Current persona: {user.persona}. {PERSONAS.get(user.persona, '')} "
        "Preserve factual meaning. Do not invent details. Return only the refined text, "
        "unless the user explicitly asks for an explanation."
    )
    output = ai_text(system, f"Task: {instruction}\n\nText:\n{source}")
    save_history(user.id, "refine", "Text refinement", source, output)
    return jsonify({"output": output})

@app.post("/api/summarize")
@jwt_required()
def summarize():
    data = request.get_json(silent=True) or {}
    source = str(data.get("text", "")).strip()
    style = str(data.get("style", "General")).strip()
    if not source:
        return fail("Paste text to summarize.")
    if len(source) > 50000:
        return fail("Text is too long. Keep it under 50,000 characters.")
    if style not in {"General", "Bullet Points", "Key Sentences"}:
        return fail("Choose General, Bullet Points, or Key Sentences.")
    style_instructions = {
        "General": "Write a concise, coherent paragraph summary.",
        "Bullet Points": "Summarize as clear bullet points. Prioritize the most important ideas.",
        "Key Sentences": "Return 3–7 key sentences that capture the central ideas."
    }
    output = ai_text(
        "You are an accurate summarization assistant. Do not add facts that are absent from the source.",
        f"{style_instructions[style]}\nIdentify the key ideas and preserve important qualifications.\n\nSource:\n{source}"
    )
    save_history(current_user().id, "summarize", f"Summary · {style}", source, output)
    return jsonify({"output": output, "style": style})

@app.post("/api/summarize-file")
@jwt_required()
def summarize_file():
    """Extract text from a supported document and summarize it using the selected style."""
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return fail("Choose a .txt, .pdf, or .docx file to summarize.")

    filename = uploaded.filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in {"txt", "pdf", "docx"}:
        return fail("Unsupported file type. Upload a .txt, .pdf, or .docx file.")

    raw = uploaded.read()
    if not raw:
        return fail("The selected file is empty.")
    if len(raw) > 5 * 1024 * 1024:
        return fail("File is too large. Maximum upload size is 5 MB.")

    try:
        if extension == "txt":
            source = raw.decode("utf-8-sig")
        elif extension == "pdf":
            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted:
                return fail("Password-protected PDFs are not supported.")
            source = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(BytesIO(raw))
            source = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                source += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    except (UnicodeDecodeError, Exception) as exc:
        app.logger.info("Document extraction failed for %s: %s", extension, type(exc).__name__)
        return fail("Could not read this document. Check that it is a valid, uncorrupted file.")

    source = source.strip()
    if not source:
        return fail("No readable text was found. Scanned/image-only PDFs need OCR and are not supported yet.")
    if len(source) > 50000:
        source = source[:50000]

    style = str(request.form.get("style", "General")).strip()
    if style not in {"General", "Bullet Points", "Key Sentences"}:
        return fail("Choose General, Bullet Points, or Key Sentences.")
    style_instructions = {
        "General": "Write a concise, coherent paragraph summary.",
        "Bullet Points": "Summarize as clear bullet points. Prioritize the most important ideas.",
        "Key Sentences": "Return 3–7 key sentences that capture the central ideas."
    }
    output = ai_text(
        "You are an accurate summarization assistant. Do not add facts that are absent from the source.",
        f"{style_instructions[style]}\nIdentify the key ideas and preserve important qualifications.\n\nSource:\n{source}"
    )
    save_history(current_user().id, "summarize", f"Document summary · {filename[:120]}", source, output)
    return jsonify({"output": output, "style": style, "filename": filename, "extracted_characters": len(source)})


@app.post("/api/translate")
@jwt_required()
def translate():
    data = request.get_json(silent=True) or {}
    source = str(data.get("text", "")).strip()
    target = str(data.get("target_language", "Hindi")).strip()
    tone = str(data.get("tone", "Preserve original tone")).strip()
    languages = {"English", "Hindi", "Telugu", "Tamil", "Kannada", "Malayalam",
                 "Bengali", "Marathi", "Gujarati", "Punjabi", "Urdu", "French",
                 "Spanish", "German", "Japanese"}
    if not source:
        return fail("Enter text to translate.")
    if target not in languages:
        return fail("That target language is not supported yet.")
    if len(source) > 30000:
        return fail("Text is too long. Keep it under 30,000 characters.")
    output = ai_text(
        "You are a professional multilingual translator. Translate naturally and contextually. "
        "Preserve meaning, names, formatting, and intent. Do not explain unless requested.",
        f"Translate the following text into {target}. Tone instruction: {tone}.\n"
        f"Use natural, fluent {target}; preserve the original meaning.\n\nText:\n{source}"
    )
    save_history(current_user().id, "translate", f"Translation · {target}", source, output)
    return jsonify({"output": output, "target_language": target})

@app.post("/api/voice/process")
@jwt_required()
def voice_process():
    data = request.get_json(silent=True) or {}
    source = str(data.get("text", "")).strip()
    action = str(data.get("action", "refine")).strip()
    target = str(data.get("target_language", "Hindi")).strip()
    if not source:
        return fail("No speech text was received. Please try recording again.")
    if action == "translate":
        prompt = f"Translate this text into {target}, preserving meaning and natural tone:\n\n{source}"
        system = "You are a careful multilingual translator. Return only the translation."
        kind, title = "voice", f"Voice translation · {target}"
    else:
        user = current_user()
        prompt = f"Correct grammar and improve clarity while preserving meaning:\n\n{source}"
        system = f"You are LinguaSense. Persona: {user.persona}. {PERSONAS.get(user.persona, '')} Return only the refined text."
        kind, title = "voice", "Voice text refinement"
    output = ai_text(system, prompt)
    save_history(current_user().id, kind, title, source, output)
    return jsonify({"output": output})

@app.get("/api/history")
@jwt_required()
def history():
    rows = (History.query.filter_by(user_id=current_user().id)
            .order_by(History.created_at.desc()).limit(100).all())
    return jsonify({"items": [{
        "id": x.id, "kind": x.kind, "title": x.title,
        "input_text": x.input_text, "output_text": x.output_text,
        "created_at": x.created_at.isoformat() + "Z"
    } for x in rows]})

@app.delete("/api/history/<int:item_id>")
@jwt_required()
def delete_history(item_id):
    row = History.query.filter_by(id=item_id, user_id=current_user().id).first()
    if not row:
        return fail("History item not found.", 404)
    db.session.delete(row)
    db.session.commit()
    return jsonify({"message": "History item deleted."})

@app.delete("/api/history")
@jwt_required()
def clear_history():
    History.query.filter_by(user_id=current_user().id).delete()
    db.session.commit()
    return jsonify({"message": "History cleared."})

@app.errorhandler(413)
def too_large(_):
    return jsonify({"error": "Request too large. Keep uploaded files under 5 MB."}), 413

@app.after_request
def add_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), geolocation=()")
    return response

@app.errorhandler(404)
def not_found(_):
    return jsonify({"error": "Endpoint not found."}), 404

@app.errorhandler(500)
def server_error(_):
    db.session.rollback()
    return jsonify({"error": "Unexpected server error. Check backend logs."}), 500

with app.app_context():
    db.create_all()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=os.getenv("FLASK_DEBUG", "0") == "1")
